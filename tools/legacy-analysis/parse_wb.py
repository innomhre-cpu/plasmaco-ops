"""Parse PLASMACO 'Finished Goods & Raw Material Stock' workbooks into normalized records.
Usage: python3 -I parse_wb.py <workbook.xlsm> <out.json>
"""
import sys, json, re, datetime, openpyxl
from openpyxl.utils import get_column_letter as L

MONTHS = ['jan','feb','mar','apr','may','jun','july','aug','sep','oct','nov','dec']

def s(v):
    if isinstance(v,(datetime.datetime,datetime.date)): return v.isoformat()
    return v

def classify(f):
    if f is None: return 'blank'
    if isinstance(f,(int,float)): return 'literal'
    if isinstance(f,(datetime.datetime,)): return 'date'
    if not isinstance(f,str): return type(f).__name__
    if not f.startswith('='): return 'text'
    if '[' in f and ']' in f: return 'external_link'
    if 'SUMIF' in f: return 'sumif'
    if '!' in f: return 'cross_sheet'
    if re.fullmatch(r'=[\d.\s+*\-/()]+', f): return 'typed_arithmetic'
    if re.fullmatch(r'=\$?[A-Z]{1,2}\$?\d+', f): return 'cell_ref'
    return 'formula'

def parse(path):
    wbf = openpyxl.load_workbook(path, data_only=False)
    wbv = openpyxl.load_workbook(path, data_only=True)
    out = {'file': path, 'sheets': []}
    for name in wbf.sheetnames:
        if name.lower() not in MONTHS: continue
        wf, wv = wbf[name], wbv[name]
        sh = {'sheet': name, 'state': wf.sheet_state, 'blocks': [], 'rm': None}
        # --- finished goods blocks
        hdrs = []
        for r in range(1, wf.max_row+1):
            for c in (1,2,3):
                if wf.cell(r,c).value == 'Product Name':
                    hdrs.append((r,c)); break
        for bi,(r,pc) in enumerate(hdrs):
            heads = [wf.cell(r,pc+k).value for k in range(0,9)]
            # date: search row r-1 cols A..C
            dcell=None
            for c in (1,2,3):
                v = wf.cell(r-1,c).value
                if v not in (None,'') and not (isinstance(v,str) and not v.strip()):
                    dcell=(L(c)+str(r-1), v, wv.cell(r-1,c).value); break
            nxt = hdrs[bi+1][0] if bi+1 < len(hdrs) else wf.max_row+1
            rows=[]; total=None
            for rr in range(r+1, nxt-1):
                nm = wf.cell(rr,pc).value
                if nm is None:
                    for cc in range(pc, pc+9):
                        v = wf.cell(rr,cc).value
                        if isinstance(v,str) and v.strip().upper().startswith('TOTAL'):
                            total = {'row': rr, 'label_cell': L(cc)+str(rr), 'formula': wf.cell(rr,cc+1).value, 'value': wv.cell(rr,cc+1).value}
                    continue
                rec = {'row': rr, 'name': nm}
                for k,key in enumerate(['open','inward','outward','closing','kg_per_m','kg','min_qty','alert'], start=1):
                    rec[key+'_f'] = s(wf.cell(rr,pc+k).value)
                    rec[key] = s(wv.cell(rr,pc+k).value)
                rows.append(rec)
            sh['blocks'].append({'header_row': r, 'product_col': L(pc), 'headers': heads, 'date_cell': dcell and [dcell[0], s(dcell[1]), s(dcell[2])], 'rows': rows, 'total': total})
        # --- raw material section
        rc=None
        for r in range(4,9):
            for c in range(10, 25):
                v = wf.cell(r,c).value
                if isinstance(v,str) and v.strip().upper()=='RECEIVED': rc=(r,c); break
            if rc: break
        if rc:
            r0,c0 = rc; rm=[]
            for rr in range(r0+1, r0+45):
                vals = [wf.cell(rr,c0+k).value for k in (-2,-1,0,1,2)]
                if all(v is None for v in vals): continue
                rm.append({'row': rr, 'date_f': s(vals[0]), 'date': s(wv.cell(rr,c0-2).value),
                           'open_f': s(vals[1]), 'open': s(wv.cell(rr,c0-1).value),
                           'recv_f': s(vals[2]), 'recv': s(wv.cell(rr,c0).value),
                           'issue_f': s(vals[3]), 'issue': s(wv.cell(rr,c0+1).value),
                           'bal_f': s(vals[4]), 'bal': s(wv.cell(rr,c0+2).value)})
            sh['rm'] = {'header_cell': L(c0)+str(r0), 'cols': [L(c0+k) for k in (-2,-1,0,1,2)], 'rows': rm}
        out['sheets'].append(sh)
    return out

if __name__ == '__main__':
    d = parse(sys.argv[1])
    json.dump(d, open(sys.argv[2],'w'), default=str)
    for sh in d['sheets']:
        print(sh['sheet'], sh['state'], 'blocks=',len(sh['blocks']), 'rm_rows=', len(sh['rm']['rows']) if sh['rm'] else None,
              'products/block=', sorted({len(b['rows']) for b in sh['blocks']}))
